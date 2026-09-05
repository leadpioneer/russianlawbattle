"""Загрузка материалов дела: ``case_files/`` + ``case_context.md``.

Поддерживаемые форматы: PDF (pypdf), DOCX (python-docx), TXT/MD (обычное
чтение). Каждый фрагмент помечается источником, фрагменты склеиваются в
единый текстовый контекст. Если грубая оценка числа токенов превышает
``Config.max_context_tokens``, текст документов суммаризируется отдельным
LLM-вызовом (модель судьи — нейтральная роль), а при необходимости —
обрезается. RAG/эмбеддинги сознательно не используются (MVP).
"""

from __future__ import annotations

import logging
import math
from dataclasses import dataclass
from pathlib import Path

from pypdf import PdfReader

from .config import Config
from .llm_client import chat

logger = logging.getLogger(__name__)

#: Поддерживаемые расширения файлов материалов дела.
SUPPORTED_EXTENSIONS: frozenset[str] = frozenset({".pdf", ".docx", ".txt", ".md"})

#: Грубая оценка числа токенов: ~3 символа на токен (смешанный рус/англ текст).
CHARS_PER_TOKEN = 3

#: Резерв токенов под промпты и ответ модели при расчёте бюджета суммаризации.
SUMMARY_RESERVE_TOKENS = 2000

#: Имя файла с промпт-контекстом дела.
CASE_CONTEXT_FILENAME = "case_context.md"

#: Минимальный бюджет (в токенах) для суммаризации, чтобы не уйти в ноль.
MIN_SUMMARY_BUDGET_TOKENS = 1000


class CaseLoadError(Exception):
    """Материалы дела не найдены, пусты или не читаются."""


@dataclass(frozen=True)
class DocumentFragment:
    """Фрагмент материалов дела с указанием источника."""

    source: str  # относительное имя файла, например "case_files/dogovor.txt"
    content: str


@dataclass(frozen=True)
class CaseMaterials:
    """Загруженные материалы дела, готовые к подстановке в промпты агентов."""

    context: str  # промпт-контекст из case_context.md (без изменений)
    documents: str  # материалы case_files/ с пометками источников (возможно суммированы)
    fragments: tuple[DocumentFragment, ...]
    estimated_tokens: int  # грубая оценка итогового контекста (context + documents)
    summarized: bool  # выполнялась ли LLM-суммаризация документов

    @property
    def is_empty(self) -> bool:
        """True, если нет ни контекста, ни документов."""
        return not self.context.strip() and not self.documents.strip()

    def full_context(self) -> str:
        """Единый текст для системных промптов: промпт-контекст + материалы дела."""
        parts: list[str] = []
        if self.context.strip():
            parts.append(f"### Промпт-контекст дела (case_context.md)\n{self.context.strip()}")
        if self.documents.strip():
            parts.append(f"### Материалы дела (case_files/)\n{self.documents.strip()}")
        return "\n\n".join(parts)


def estimate_tokens(text: str) -> int:
    """Оценка числа токенов без внешних токенайзеров (~CHARS_PER_TOKEN симв./токен)."""
    return math.ceil(len(text) / CHARS_PER_TOKEN)


# --- читатели отдельных форматов -------------------------------------------

def _read_txt(path: Path) -> str:
    return path.read_text(encoding="utf-8-sig")


def _read_pdf(path: Path) -> str:
    """Извлечь текст PDF постранично (pypdf)."""
    pages: list[str] = []
    for number, page in enumerate(PdfReader(str(path)).pages, start=1):
        text = (page.extract_text() or "").strip()
        pages.append(f"[страница {number}]\n{text}" if text else f"[страница {number}] (текст не извлечён)")
    return "\n\n".join(pages)


def _read_docx(path: Path) -> str:
    """Извлечь текст DOCX: абзацы + ячейки таблиц (python-docx)."""
    from docx import Document  # локальный импорт: нужен только для docx

    document = Document(str(path))
    lines = [p.text.strip() for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                lines.append(" | ".join(cells))
    return "\n".join(lines)


def read_document(path: Path) -> str:
    """Прочитать документ в текст; неподдерживаемый формат -> CaseLoadError."""
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return _read_pdf(path)
    if suffix == ".docx":
        return _read_docx(path)
    if suffix in (".txt", ".md"):
        return _read_txt(path)
    raise CaseLoadError(
        f"Неподдерживаемый формат документа: {path.name} ({suffix or 'без расширения'})."
    )


# --- сборка контекста дела ---------------------------------------------------

def _render_fragment(fragment: DocumentFragment) -> str:
    """Фрагмент с пометкой источника (попадает в промпты агентов)."""
    return f"--- Источник: {fragment.source} ---\n{fragment.content.strip()}"


def load_case_context(project_root: Path) -> str:
    """Промпт-контекст дела из case_context.md (обязателен)."""
    path = project_root / CASE_CONTEXT_FILENAME
    if not path.is_file():
        raise CaseLoadError(
            f"Не найден файл {path}\n"
            f"Скопируйте case_context.md.example в case_context.md и опишите суть дела:\n"
            f"стороны, позиции, что должно решить судье."
        )
    return path.read_text(encoding="utf-8-sig")


def load_documents(case_files_dir: Path) -> tuple[DocumentFragment, ...]:
    """Читать все поддерживаемые файлы каталога case_files/ (по алфавиту)."""
    if not case_files_dir.is_dir():
        logger.warning("Каталог %s не найден — материалы дела пусты.", case_files_dir)
        return ()

    fragments: list[DocumentFragment] = []
    for path in sorted(p for p in case_files_dir.iterdir() if p.is_file()):
        if path.name == ".gitkeep":
            continue
        if path.suffix.lower() not in SUPPORTED_EXTENSIONS:
            logger.warning("Пропущен файл неподдерживаемого формата: %s", path.name)
            continue
        text = read_document(path)
        if not text.strip():
            logger.warning("Файл пуст или текст не извлечён (скан?): %s", path.name)
            continue
        fragments.append(DocumentFragment(source=f"case_files/{path.name}", content=text))
        logger.info("Загружен документ %s: %d символов.", path.name, len(text))
    return tuple(fragments)


def summarize_documents(documents_text: str, target_chars: int) -> str:
    """Сжать материалы дела отдельным LLM-вызовом (модель судьи, temperature=0.1).

    Суммаризатор обязан сохранить факты, даты, суммы, реквизиты, стороны
    и ссылки на нормы права — без добавления того, чего нет в материалах.
    """
    system_prompt = (
        "Ты — ассистент-юрист, готовящий материалы дела для судебного процесса. "
        "Сожми предоставленные материалы, ОБЯЗАТЕЛЬНО сохранив: наименования и процессуальный "
        "статус сторон, ключевые факты и хронологию, даты, суммы, номера и реквизиты "
        "документов, важные условия договоров и ссылки на нормы права. "
        "Ничего не выдумывай и не добавляй фактов, которых нет в материалах. "
        "Ответ — только сжатый текст материалов, без комментариев."
    )
    user_prompt = (
        f"Сожми материалы дела примерно до {target_chars} символов "
        f"(~{target_chars // CHARS_PER_TOKEN} токенов).\n\nМАТЕРИАЛЫ ДЕЛА:\n{documents_text}"
    )
    return chat(
        "judge",
        system_prompt,
        [{"role": "user", "content": user_prompt}],
        temperature=0.1,
    )


def _truncate(text: str, max_chars: int) -> str:
    """Обрезать текст по границе абзаца с пометкой об обрезке."""
    if len(text) <= max_chars:
        return text
    cut = text[:max_chars].rsplit("\n", 1)[0]
    return cut + "\n\n[... текст обрезан из-за лимита контекста ...]"


def load_case(cfg: Config) -> CaseMaterials:
    """Загрузить материалы дела: case_context.md + case_files/ с пометками источников.

    Если грубая оценка контекста превышает ``cfg.max_context_tokens``:
    1) документы суммаризируются LLM (модель судьи);
    2) если и после этого не влезает — обрезаются по границе абзаца.

    :raises CaseLoadError: если нет case_context.md.
    """
    context = load_case_context(cfg.project_root)
    fragments = load_documents(cfg.case_files_dir)
    documents = "\n\n".join(_render_fragment(f) for f in fragments)

    estimated = estimate_tokens(context) + estimate_tokens(documents)
    summarized = False

    if estimated > cfg.max_context_tokens and documents.strip():
        budget = (
            cfg.max_context_tokens
            - estimate_tokens(context)
            - SUMMARY_RESERVE_TOKENS
        )
        target_chars = max(budget, MIN_SUMMARY_BUDGET_TOKENS) * CHARS_PER_TOKEN
        logger.info(
            "Контекст дела ~%d токенов превышает лимит %d — суммаризирую документы до ~%d символов.",
            estimated,
            cfg.max_context_tokens,
            target_chars,
        )
        documents = summarize_documents(documents, target_chars)
        summarized = True

    # Последняя инстанция: гарантировать влезание даже после суммаризации.
    if estimate_tokens(documents) + estimate_tokens(context) > cfg.max_context_tokens:
        allowed_chars = max(
            cfg.max_context_tokens - estimate_tokens(context), MIN_SUMMARY_BUDGET_TOKENS
        ) * CHARS_PER_TOKEN
        documents = _truncate(documents, allowed_chars)
        estimated = estimate_tokens(context) + estimate_tokens(documents)

    estimated = estimate_tokens(context) + estimate_tokens(documents)
    logger.info(
        "Материалы дела готовы: документов=%d, суммаризация=%s, оценка контекста ~%d токенов.",
        len(fragments),
        summarized,
        estimated,
    )
    return CaseMaterials(
        context=context,
        documents=documents,
        fragments=fragments,
        estimated_tokens=estimated,
        summarized=summarized,
    )
